from fastapi import FastAPI, Request, HTTPException, Depends, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc
from typing import Optional, List
from pydantic import BaseModel
import httpx
import os
import io
import docx
import fitz
import base64
import pandas as pd
import json
import re
import zipfile
from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage, SystemMessage

# 🌟 导入同目录下的数据库和 agent 逻辑
from database import get_db, init_db, ChatHistory, SessionLocal, QuizRecord, StudentScore, User
from agent_service import A25Agent

# ==========================================
# 1. 环境配置与应用初始化
# ==========================================
app = FastAPI(title="A25 Agent Gateway ")
from fastapi import Response  # 确保导入了 Response

#gemini辅助生成
@app.middleware("http")
async def disable_x_frame_options(request: Request, call_next):
    response: Response = await call_next(request)

    if "X-Frame-Options" in response.headers:
        del response.headers["X-Frame-Options"]
    #  2：显式允许所有站点嵌套（兼容老浏览器）
    response.headers["X-Frame-Options"] = "ALLOWALL"

    # 3：设置现代化的 CSP 策略，允许被学习通等域名嵌套
    # frame-ancestors * 代表允许被任何站点嵌套，调试期建议这样，后期可改为具体域名
    response.headers["Content-Security-Policy"] = "frame-ancestors *;"

    return response

# 允许跨域
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

agent = A25Agent()

# 微服务内网地址
RAG_SERVICE_URL = os.getenv("RAG_SERVICE_URL", "http://rag-service:8001")
GRADING_SERVICE_URL = os.getenv("GRADING_SERVICE_URL", "http://grading-service:8002")

# 大模型 API 配置
DASHSCOPE_BASE_URL = os.getenv("DASHSCOPE_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1")
DASHSCOPE_API_KEY = os.getenv("DASHSCOPE_API_KEY") or os.getenv("OPENAI_API_KEY")

@app.on_event("startup")
async def startup_event():
    init_db()
    print("🚀 [网关] A25 后端云原生引擎已成功点火！运行环境: Docker/Linux")

@app.get("/")
async def root():
    return {"message": "A25 Gateway is running!"}

# ==========================================
# 2. 数据模型区 (Pydantic Models)
# ==========================================
class ChatMessage(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    user_id: str = "student_001"
    course_id: str = "default_course"
    platform: str = "web"
    messages: List[ChatMessage]
    page_context: Optional[dict] = None

class QuizRecordCreate(BaseModel):
    user_id: str
    course_id: str
    question: str
    options: str
    user_choice: str
    correct_answer: str
    is_correct: bool
    explanation: str = ""

class HomeworkRequest(BaseModel):
    user_id: str
    course_id: str
    question_text: str
    content_type: str
    content: str

class ExportDocRequest(BaseModel):
    content: str

class TargetedPracticeRequest(BaseModel):
    course_id: str
    error_context: str
    suggestion: str = ""

class PromptEnhanceRequest(BaseModel):
    short_desc: str
    role: str

class LoginRequest(BaseModel):
    account_id: str
    password: str
    role: str

class RegisterRequest(BaseModel):
    account_id: str
    password: str
    role: str
    name: str

# ==========================================
# 3. 核心大模型路由区 (Chat & Agent)
# ==========================================
@app.post("/api/chat")
async def chat_endpoint(request: ChatRequest): # 🌟 回归你原本优秀的 Pydantic 模型拦截！
    if not request.messages:
        raise HTTPException(status_code=400, detail="Messages cannot be empty")

    # 1. 精准提取当前消息和前端发来的上下文 (原汁原味的旧代码逻辑)
    current_message = request.messages[-1].content
    live_context = [{"role": msg.role, "content": msg.content} for msg in request.messages[:-1]]

    # 2. 读取数据库历史记录
    db_read = SessionLocal()
    history_context = []
    user_name = "刘明剑同学"
    try:
        # 顺便查一下名字
        user_info = db_read.query(User).filter(User.account_id == request.user_id).first()
        if user_info:
            user_name = user_info.name

        recent_history = db_read.query(ChatHistory).filter(
            ChatHistory.user_id == request.user_id
        ).order_by(desc(ChatHistory.id)).limit(5).all()
        recent_history.reverse()
        history_context = [{"role": h.role, "content": h.content} for h in recent_history]
    finally:
        db_read.close()

    # 将数据库历史与当前页面活体上下文合并
    full_memory = history_context + live_context

    # 3. 跨服请求 RAG 微服务 (兼容加强版)
    knowledge_context = ""
    try:
        async with httpx.AsyncClient(timeout=30.0) as client:
            rag_resp = await client.post(
                f"{RAG_SERVICE_URL}/retrieve",
                json={"query": current_message, "course_id": request.course_id}
            )
            if rag_resp.status_code == 200:
                # 🚀 无论 RAG 返回什么神仙格式，统统拿下！
                try:
                    rag_data = rag_resp.json()
                    if isinstance(rag_data, dict):
                        # 盲猜覆盖常见的所有 Key
                        knowledge_context = rag_data.get("data", rag_data.get("text", rag_data.get("context",
                                                                                                   rag_data.get(
                                                                                                       "answer", ""))))
                    elif isinstance(rag_data, str):
                        # 如果 RAG 直接返回了纯字符串
                        knowledge_context = rag_data
                except Exception:
                    # 如果连 JSON 都不是，直接拿原生纯文本响应体
                    knowledge_context = rag_resp.text

                if knowledge_context:
                    print(f"✅ [网关] 成功从 RAG 提取到上下文资料，长度: {len(knowledge_context)} 字符")
                else:
                    print(f"⚠️ [网关] RAG 返回了 200，但解析不到文本内容。原始数据: {rag_resp.text}")
            else:
                print(f"⚠️ [网关] RAG 请求异常，状态码: {rag_resp.status_code}")
    except Exception as e:
        print(f"⚠️ [网关] RAG 网络请求彻底崩溃: {type(e).__name__} - {str(e)}")

    # 4. 流式生成与保存 (完美继承你的旧逻辑)
    async def stream_and_save():
        full_ai_response = ""
        try:
            async for chunk in agent.execute_stream(
                user_input=current_message,
                chat_history=full_memory,      # 🌟 传入完美拼接的历史
                knowledge_context=knowledge_context,
                page_context=request.page_context,
                user_id=request.user_id,
                course_id=request.course_id
            ):
                full_ai_response += chunk
                yield chunk

            # 生成完毕后保存入库 (使用你原本的开闭 Session 逻辑，最稳妥)
            if full_ai_response.strip():
                db_write = SessionLocal()
                try:
                    # ⚠️ 注意：这里角色存 'assistant' 以匹配前端的渲染判断
                    db_write.add(ChatHistory(user_id=request.user_id, role="user", content=current_message))
                    db_write.add(ChatHistory(user_id=request.user_id, role="assistant", content=full_ai_response))
                    db_write.commit()
                except Exception as e:
                    db_write.rollback()
                    print(f"❌ [网关] 记录保存失败: {e}")
                finally:
                    db_write.close()

        except Exception as e:
            # 🛡️ 终极防爆盾：哪怕大模型真崩了，也能安全吐出报错，不炸断 ASGI
            error_msg = f"\n\n[系统状态] ⚠️ 大模型响应中断: {str(e)}"
            print(error_msg)
            yield error_msg

    return StreamingResponse(stream_and_save(), media_type="text/event-stream")

# ==========================================
# 4. 账号鉴权路由区
# ==========================================
@app.post("/api/register")
async def register(request: RegisterRequest, db: Session = Depends(get_db)):
    existing = db.query(User).filter(User.account_id == request.account_id).first()
    if existing:
        raise HTTPException(status_code=400, detail="账号已存在")
    new_user = User(account_id=request.account_id, password=request.password, role=request.role, name=request.name)
    db.add(new_user)
    db.commit()
    return {"status": "success", "message": "注册成功"}

@app.post("/api/login")
async def login(request: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.account_id == request.account_id, User.role == request.role).first()
    if not user or user.password != request.password:
        raise HTTPException(status_code=401, detail="账号或密码错误")
    return {"status": "success", "user": {"account_id": user.account_id, "role": user.role, "name": user.name}}


# ==========================================
# 5. 原生业务路由区 (成绩、作业、错题本等)
# ==========================================
@app.post("/api/quiz")
async def generate_quiz_endpoint(request: Request):
    data = await request.json()
    page_context = data.get("page_context")
    quiz_result = await agent.generate_quiz(page_context)
    if "error" in quiz_result: return JSONResponse(status_code=400, content=quiz_result)
    return quiz_result

@app.post("/api/save_quiz_record")
async def save_quiz_record(record: QuizRecordCreate, db: Session = Depends(get_db)):
    try:
        new_record = QuizRecord(**record.dict())
        db.add(new_record)
        db.commit()
        db.refresh(new_record)
        return {"status": "success", "id": new_record.id}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")

@app.post("/api/grade_homework")
async def grade_homework_endpoint(request: HomeworkRequest):
    """带 Word 原生侧边栏批注的作业批改引擎"""
    print(f"📝 [作业批改] 收到作业，类型: {request.content_type}")
    actual_content = request.content
    doc_obj = None
    annotated_file_b64 = None

    if request.content_type == "docx":
        base64_data = request.content.split(",")[-1] if "," in request.content else request.content
        file_bytes = base64.b64decode(base64_data)
        try:
            doc_obj = docx.Document(io.BytesIO(file_bytes))
            actual_content = "\n".join([para.text for para in doc_obj.paragraphs])
            print(" 成功从上传的 Word 中解剖出纯文本。")
        except zipfile.BadZipFile:
            print(" 拦截：检测到伪造的 .docx 或文件已损坏！")
            return {"status": "error", "message": "文件格式错误！请确保上传的是真实的 .docx 文件。"}
        except Exception as e:
            return {"status": "error", "message": f"文档读取失败: {e}"}
    elif request.content_type == "doc":
        return {"status": "error", "message": "⚠️ 格式不支持：仅支持 .docx 格式。请用 Word 另存为 .docx 后重新上传！"}

    # 调用 AI 批改
    result_str = await agent.grade_homework(
        question_text=request.question_text,
        content_type=request.content_type,
        content=actual_content,
        course_id=request.course_id
    )

    if request.content_type == "docx" and doc_obj:
        try:
            grading_data = json.loads(result_str)
            def clean_str(s): return re.sub(r'\W+', '', s)

            for ann in grading_data.get("annotations", []):
                target_text = ann.get("target_text", "") or ann.get("quoted_text", "")
                if not target_text: continue
                clean_target = clean_str(target_text)

                found_paragraph = None
                max_overlap = 0

                for paragraph in doc_obj.paragraphs:
                    if not paragraph.text.strip() or not paragraph.runs: continue
                    clean_para = clean_str(paragraph.text)

                    if clean_target and clean_target in clean_para:
                        found_paragraph = paragraph
                        break

                    elif clean_target and clean_para in clean_target:
                        if len(clean_para) > 10 and len(clean_para) > max_overlap:
                            max_overlap = len(clean_para)
                            found_paragraph = paragraph

                error_type = ann.get('error_type', '逻辑错误')
                explanation = ann.get('explanation', '解答存在瑕疵。')
                suggestion = ann.get('suggestion', '请重新检查此处的上下文逻辑。')
                location = ann.get('location', '未知位置')

                comment_text = f"【{error_type}】\n{explanation}\n💡建议：{suggestion}"

                if found_paragraph:
                    doc_obj.add_comment(found_paragraph.runs, text=comment_text, author="A25教授", initials="AI")
                else:
                    for p in reversed(doc_obj.paragraphs):
                        if p.text.strip() and p.runs:
                            doc_obj.add_comment(p.runs,
                                                text=f"【原位置：{location}】\n(原文未精确定位: {target_text})\n{comment_text}",
                                                author="A25教授", initials="AI")
                            break

            output_stream = io.BytesIO()
            doc_obj.save(output_stream)
            annotated_file_b64 = "data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64," + base64.b64encode(
                output_stream.getvalue()).decode("utf-8")
            print(" 成功将 AI 批注写入原生 Word 侧边栏！")
        except Exception as e:
            print(f"⚠ 写入 Word 批注失败: {e}")

    return {"status": "success", "grading_result": result_str, "annotated_file": annotated_file_b64}#chatgpt辅助修改返回格式

@app.post("/api/parse_question_file")
async def parse_question_file(file: UploadFile = File(...)):
    print(f"📄 [文档解析] 收到题目文件: {file.filename}")
    try:
        content = await file.read()
        extracted_text = ""

        if file.filename.lower().endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            extracted_text = "\n".join([para.text for para in doc.paragraphs])
        elif file.filename.lower().endswith(".pdf"):
            pdf_document = fitz.open(stream=content, filetype="pdf")
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                extracted_text += page.get_text() + "\n"

            if len(extracted_text.strip()) < 20:
                print("🚨 检测到纯图片/扫描版 PDF！启动 OCR 视觉大模型...")
                extracted_text = ""
                vision_llm = ChatOpenAI(
                    base_url=DASHSCOPE_BASE_URL,
                    model="qwen-vl-max",
                    temperature=0,
                    api_key=DASHSCOPE_API_KEY,
                )
                for page_num in range(len(pdf_document)):
                    page = pdf_document.load_page(page_num)
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                    img_data = pix.tobytes("jpeg")
                    base64_img = base64.b64encode(img_data).decode('utf-8')
                    try:
                        msg = HumanMessage(content=[
                            {"type": "text", "text": "提取文字内容。"},
                            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_img}"}}
                        ])
                        resp = await vision_llm.ainvoke([msg])
                        extracted_text += resp.content + "\n\n"
                    except Exception as e:
                        pass
            pdf_document.close()
        else:
            return JSONResponse(status_code=400, content={"error": "仅支持 .docx 和 .pdf"})
        return {"status": "success", "text": extracted_text.strip()}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": f"解析失败: {str(e)}"})

@app.post("/api/analyze_scores")
async def analyze_scores_endpoint(file: UploadFile = File(...), db: Session = Depends(get_db)):
    try:
        content = await file.read()
        if file.filename.lower().endswith(".csv"):
            try: df = pd.read_csv(io.BytesIO(content), encoding='utf-8')
            except: df = pd.read_csv(io.BytesIO(content), encoding='gbk')
        else: df = pd.read_excel(io.BytesIO(content))

        df.dropna(axis=1, how='all', inplace=True)
        df = df.fillna(0)

        if "学号" in df.columns:
            for row in df.to_dict(orient="records"):
                uid = str(row.get("学号", "")).strip()
                if not uid or uid in ["0.0", "0"]: continue
                s_name = str(row.get("姓名", "未知"))
                s_data = {k: float(v) for k, v in row.items() if k.strip() not in ["学号", "姓名", "序号", "id", "总分", "平均", "排名", "优等生"] and isinstance(v, (int, float))}

                existing = db.query(StudentScore).filter(StudentScore.user_id == uid).first()
                if existing:
                    existing.score_data = s_data
                    existing.student_name = s_name
                else: db.add(StudentScore(user_id=uid, student_name=s_name, score_data=s_data))
            db.commit()

        numeric_cols = df.select_dtypes(include=['number', 'float', 'int']).columns.tolist()
        blacklist = ["学号", "序号", "id", "总分", "平均", "排名", "优等生"]
        knowledge_points = [col for col in numeric_cols if col.strip() not in blacklist]

        radar_data = []
        for kp in knowledge_points:
            radar_data.append({"subject": kp, "A": round(df[kp].mean(), 1), "fullMark": 100})

        weakest_subject = min(radar_data, key=lambda x: x["A"])["subject"] if radar_data else "未知"
        bar_data = []
        if weakest_subject in df.columns:
            scores = df[weakest_subject]
            bar_data = [
                {"name": "0-20分", "人数": int(((scores >= 0) & (scores < 20)).sum())},
                {"name": "20-40分", "人数": int(((scores >= 20) & (scores < 40)).sum())},
                {"name": "40-60分", "人数": int(((scores >= 40) & (scores < 60)).sum())},
                {"name": "60-80分", "人数": int(((scores >= 60) & (scores < 80)).sum())},
                {"name": "80-100分", "人数": int((scores >= 80).sum())},
            ]

        return {"status": "success", "radar_data": radar_data, "bar_data": bar_data, "weakest_subject": weakest_subject, "total_students": len(df)}
    except Exception as e:
        db.rollback()
        return JSONResponse(status_code=500, content={"error": str(e)})

@app.get("/api/my_score")
async def get_my_score(account_id: str, db: Session = Depends(get_db)):
    student_record = db.query(StudentScore).filter(StudentScore.user_id == account_id).first()
    if not student_record or not student_record.score_data:
        return {"status": "not_found", "message": "尚未导入成绩"}
    my_radar = [{"subject": k, "A": v, "fullMark": 100} for k, v in student_record.score_data.items()]
    my_weakest = min(my_radar, key=lambda x: x["A"])["subject"] if my_radar else "暂无"
    return {"status": "success", "student_name": student_record.student_name, "radar_data": my_radar, "weakest_subject": my_weakest}

@app.post("/api/export_doc")
async def export_doc_endpoint(request: ExportDocRequest):
    try:
        doc = docx.Document()
        doc.add_heading('A25 智能教研产出物', level=0)
        for line in request.content.split('\n'):
            line = line.strip()
            if not line: continue
            if line.startswith('# '): doc.add_heading(line[2:], level=1)
            elif line.startswith('## '): doc.add_heading(line[3:], level=2)
            elif line.startswith('### '): doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '): doc.add_paragraph(line[2:], style='List Bullet')
            else: doc.add_paragraph(line)
        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)
        return StreamingResponse(output_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={'Content-Disposition': 'attachment; filename="A25_Smart_Plan.docx"'})
    except Exception as e:
        raise HTTPException(status_code=500, detail="生成文档失败")

@app.post("/api/upload_knowledge")
async def upload_knowledge_endpoint(file: UploadFile = File(...)):
    """文件通过 HTTP 转发给后端的 RAG 微服务进行入库"""
    if not file.filename.lower().endswith(".pdf"): return JSONResponse(status_code=400, content={"error": "仅支持 PDF"})
    try:
        file_content = await file.read()
        print(f"📡 [网关] 正在将文件 {file.filename} 转发给 RAG 微服务进行向量化入库...")
        async with httpx.AsyncClient() as client:
            files = {'file': (file.filename, file_content, file.content_type)}
            resp = await client.post(f"{RAG_SERVICE_URL}/api/upload", files=files, timeout=60.0)
            return resp.json()
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": f"网关转发失败: {str(e)}"})

@app.post("/api/generate_practice")
async def generate_practice(request: TargetedPracticeRequest):
    try:
        result_json_str = await agent.generate_targeted_exercises(course_id=request.course_id, error_annotation={
            "explanation": request.error_context, "suggestion": request.suggestion})
        print(f"【DEBUG】大模型原始输出: {result_json_str}")
        return {"status": "success", "data": json.loads(result_json_str)}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.post("/api/enhance_prompt")
async def enhance_prompt(request: PromptEnhanceRequest):
    system_instruction = "你是一位世界顶级的 AI 提示词工程师...直接输出结果。"
    try:
        messages = [SystemMessage(content=system_instruction), HumanMessage(content=f"扩写：{request.short_desc}")]
        response = await agent.llm.ainvoke(messages)
        return {"status": "success", "enhanced_prompt": response.content}
    except Exception as e:
        return {"status": "error", "message": str(e)}

import asyncio
import random
from fastapi import Request

@app.post("/api/mock_ai_generate")
async def mock_ai_generate(request: Request):
    """
    假装是大模型的挡板接口。
    """
    # 模拟大模型思考的延迟A
    await asyncio.sleep(random.uniform(1.0, 2.5))
    return {
        "status": "success",
        "data": {
            "question": "这是一道极限抗压测试题",
            "answer": "A"
        }
    }