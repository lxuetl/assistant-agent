from fastapi import FastAPI, HTTPException, Depends, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import desc
from typing import Optional, List
from fastapi import UploadFile, File
import docx
from docx.shared import RGBColor
import fitz
import io
import pandas as pd
import json
import re
import aiofiles
import os
import base64
import zipfile # 🌟 核心新增：用于捕捉虚假的 docx 文件
from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage, SystemMessage
from pydantic import BaseModel
from agent_service import A25Agent
from database import get_db, QuizRecord, ChatHistory, StudentScore, SessionLocal, engine, Base, User
from vector_store import build_vector_store

#import win32com.client
#import pythoncom
import tempfile

# 1. 初始化应用和 Agent
app = FastAPI(title="A25 Agent API")
Base.metadata.create_all(bind=engine)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

agent = A25Agent()
print("🚀 我是后端，我活得好好的！")

# ==========================================
# 🌟 数据模型区
# ==========================================
class ChatMessage(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    user_id: str = "student_001"
    course_id: str = "default_course"
    platform: str = "web"
    messages: List[ChatMessage]
    page_context: Optional[dict] = None

class QuizRecordCreate(BaseModel):
    user_id: str
    course_id: str
    question: str
    options: str
    user_choice: str
    correct_answer: str
    is_correct: bool
    explanation: str = ""

class HomeworkRequest(BaseModel):
    user_id: str
    course_id: str
    question_text: str
    content_type: str
    content: str

class ExportDocRequest(BaseModel):
    content: str

class TargetedPracticeRequest(BaseModel):
    course_id: str
    error_context: str
    suggestion: str = ""

class PromptEnhanceRequest(BaseModel):
    short_desc: str
    role: str

class LoginRequest(BaseModel):
    account_id: str
    password: str
    role: str

class RegisterRequest(BaseModel):
    account_id: str
    password: str
    role: str  # student 或 teacher
    name: str

# ==========================================
# 🌟 登录接口
# ==========================================
@app.post("/api/login")
async def login(request: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(
        User.account_id == request.account_id,
        User.role == request.role
    ).first()
    
    if not user or user.password != request.password:  # 实际用哈希验证
        raise HTTPException(status_code=401, detail="账号或密码错误")
    
    return {
        "status": "success",
        "user": {
            "account_id": user.account_id,
            "role": user.role,
            "name": user.name
        }
    }

# ==========================================
# 🌟 API 路由区
# ==========================================
@app.post("/api/chat")
async def chat_endpoint(request: ChatRequest):
    if not request.messages:
        raise HTTPException(status_code=400, detail="Messages cannot be empty")
    current_message = request.messages[-1].content
    live_context = [{"role": msg.role, "content": msg.content} for msg in request.messages[:-1]]
    db_read = SessionLocal()
    history_context = []
    try:
        recent_history = db_read.query(ChatHistory).filter(ChatHistory.user_id == request.user_id).order_by(desc(ChatHistory.id)).limit(5).all()
        recent_history.reverse()
        history_context = [{"role": h.role, "content": h.content} for h in recent_history]
    finally:
        db_read.close()

    full_memory = history_context + live_context

    async def stream_and_save():
        full_ai_response = ""
        async for chunk in agent.execute_stream(user_input=current_message, page_context=request.page_context, chat_history=full_memory, user_id=request.user_id, course_id=request.course_id):
            full_ai_response += chunk
            yield chunk

        db_write = SessionLocal()
        try:
            db_write.add(ChatHistory(user_id=request.user_id, role="user", content=current_message))
            db_write.add(ChatHistory(user_id=request.user_id, role="ai", content=full_ai_response))
            db_write.commit()
        except Exception as e:
            db_write.rollback()
        finally:
            db_write.close()
    return StreamingResponse(stream_and_save(), media_type="text/event-stream")

@app.post("/api/quiz")
async def generate_quiz_endpoint(request: Request):
    data = await request.json()
    page_context = data.get("page_context")
    quiz_result = await agent.generate_quiz(page_context)
    if "error" in quiz_result:
        return JSONResponse(status_code=400, content=quiz_result)
    return quiz_result

@app.post("/api/save_quiz_record")
async def save_quiz_record(record: QuizRecordCreate, db: Session = Depends(get_db)):
    try:
        new_record = QuizRecord(user_id=record.user_id, course_id=record.course_id, question=record.question, options=record.options, user_choice=record.user_choice, correct_answer=record.correct_answer, is_correct=record.is_correct, explanation=record.explanation)
        db.add(new_record)
        db.commit()
        db.refresh(new_record)
        return {"status": "success", "id": new_record.id}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")


# 🌟 恢复：完美的【原生侧边栏批注】与假文件防御机制！
@app.post("/api/grade_homework")
async def grade_homework_endpoint(request: HomeworkRequest):
    print(f"📝 [作业批改] 收到作业，类型: {request.content_type}")
    actual_content = request.content
    doc_obj = None
    annotated_file_b64 = None

    # 1. 精准读取并拦截非法 docx 格式
    if request.content_type == "docx":
        base64_data = request.content.split(",")[-1] if "," in request.content else request.content
        file_bytes = base64.b64decode(base64_data)
        try:
            doc_obj = docx.Document(io.BytesIO(file_bytes))
            actual_content = "\n".join([para.text for para in doc_obj.paragraphs])
            print("📄 成功从上传的 Word 中解剖出纯文本。")
        except zipfile.BadZipFile:
            print("❌ 拦截：检测到伪造的 .docx 或文件已损坏！")
            return {"status": "error", "message": "文件格式错误！请确保上传的是真实的 .docx 文件，切勿直接修改后缀名。"}
        except Exception as e:
            return {"status": "error", "message": f"文档读取失败: {e}"}

    # 如果是老旧的 doc，直接让用户去转存
    elif request.content_type == "doc":
        return {"status": "error",
                "message": "⚠️ 格式不支持：系统已升级为原生批注机制，仅支持 .docx 格式。请用 Word 另存为 .docx 后重新上传！"}

    # 2. 调用 AI 批改引擎
    result_str = await agent.grade_homework(
        question_text=request.question_text,
        content_type=request.content_type,
        content=actual_content,
        course_id=request.course_id
    )

    # 3. 🌟 完美的【原生侧边栏批注】逻辑
    if request.content_type == "docx" and doc_obj:
        try:
            grading_data = json.loads(result_str)

            def clean_str(s):
                return re.sub(r'\W+', '', s)

            for ann in grading_data.get("annotations", []):
                target_text = ann.get("quoted_text", "")
                if not target_text: continue

                clean_target = clean_str(target_text)
                found_paragraph = None

                for paragraph in doc_obj.paragraphs:
                    if not paragraph.text.strip() or not paragraph.runs:
                        continue
                    clean_para = clean_str(paragraph.text)
                    if clean_target and (clean_target in clean_para or clean_para in clean_target):
                        found_paragraph = paragraph
                        break

                comment_text = f"【{ann['error_type']}】\n{ann['explanation']}\n💡建议：{ann['suggestion']}"

                # 🌟 执行强大的原生批注
                if found_paragraph:
                    doc_obj.add_comment(found_paragraph.runs, text=comment_text, author="A25教授", initials="AI")
                else:
                    for p in doc_obj.paragraphs:
                        if p.text.strip() and p.runs:
                            doc_obj.add_comment(p.runs, text=f"【原位置：{ann['location']}】\n" + comment_text,
                                                author="A25教授", initials="AI")
                            break

            output_stream = io.BytesIO()
            doc_obj.save(output_stream)
            annotated_file_b64 = "data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64," + base64.b64encode(
                output_stream.getvalue()).decode("utf-8")
            print("🎉 成功将 AI 批注写入原生 Word 侧边栏！")

        except Exception as e:
            print(f"⚠️ 写入 Word 批注失败: {e}")

    return {
        "status": "success",
        "grading_result": result_str,
        "annotated_file": annotated_file_b64
    }


@app.post("/api/parse_question_file")
async def parse_question_file(file: UploadFile = File(...)):
    print(f"📄 [文档解析] 收到题目文件: {file.filename}")
    try:
        content = await file.read()
        extracted_text = ""

        if file.filename.lower().endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            extracted_text = "\n".join([para.text for para in doc.paragraphs])
        elif file.filename.lower().endswith(".pdf"):
            pdf_document = fitz.open(stream=content, filetype="pdf")
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                extracted_text += page.get_text() + "\n"

            if len(extracted_text.strip()) < 20:
                print("🚨 检测到纯图片/扫描版 PDF！启动 OCR...")
                extracted_text = ""
                vision_llm = ChatOpenAI(base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
                                        model="qwen-vl-max", temperature=0)
                for page_num in range(len(pdf_document)):
                    page = pdf_document.load_page(page_num)
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                    img_data = pix.tobytes("jpeg")
                    base64_img = base64.b64encode(img_data).decode('utf-8')
                    try:
                        msg = HumanMessage(content=[
                            {"type": "text", "text": "提取文字内容。"},
                            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_img}"}}
                        ])
                        resp = await vision_llm.ainvoke([msg])
                        extracted_text += resp.content + "\n\n"
                    except Exception as e:
                        pass
            pdf_document.close()
        else:
            return JSONResponse(status_code=400, content={"error": "仅支持 .docx 和 .pdf"})
        return {"status": "success", "text": extracted_text.strip()}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": f"解析失败: {str(e)}"})


@app.post("/api/analyze_scores")
async def analyze_scores_endpoint(file: UploadFile = File(...), db: Session = Depends(get_db)):
    try:
        content = await file.read()
        if file.filename.lower().endswith(".csv"):
            try:
                df = pd.read_csv(io.BytesIO(content), encoding='utf-8')
            except:
                df = pd.read_csv(io.BytesIO(content), encoding='gbk')
        else:
            df = pd.read_excel(io.BytesIO(content))

        df.dropna(axis=1, how='all', inplace=True)
        df = df.fillna(0)

        if "学号" in df.columns:
            for row in df.to_dict(orient="records"):
                uid = str(row.get("学号", "")).strip()
                if not uid or uid in ["0.0", "0"]: continue
                s_name = str(row.get("姓名", "未知"))
                s_data = {k: float(v) for k, v in row.items() if
                          k.strip() not in ["学号", "姓名", "序号", "id", "总分", "平均", "排名",
                                            "优等生"] and isinstance(v, (int, float))}

                existing = db.query(StudentScore).filter(StudentScore.user_id == uid).first()
                if existing:
                    existing.score_data = s_data
                    existing.student_name = s_name
                else:
                    db.add(StudentScore(user_id=uid, student_name=s_name, score_data=s_data))
            db.commit()

        numeric_cols = df.select_dtypes(include=['number', 'float', 'int']).columns.tolist()
        blacklist = ["学号", "序号", "id", "总分", "平均", "排名", "优等生"]
        knowledge_points = [col for col in numeric_cols if col.strip() not in blacklist]

        radar_data = []
        for kp in knowledge_points:
            radar_data.append({"subject": kp, "A": round(df[kp].mean(), 1), "fullMark": 100})

        weakest_subject = min(radar_data, key=lambda x: x["A"])["subject"] if radar_data else "未知"
        bar_data = []
        if weakest_subject in df.columns:
            scores = df[weakest_subject]
            bar_data = [
                {"name": "0-20分", "人数": int(((scores >= 0) & (scores < 20)).sum())},
                {"name": "20-40分", "人数": int(((scores >= 20) & (scores < 40)).sum())},
                {"name": "40-60分", "人数": int(((scores >= 40) & (scores < 60)).sum())},
                {"name": "60-80分", "人数": int(((scores >= 60) & (scores < 80)).sum())},
                {"name": "80-100分", "人数": int((scores >= 80).sum())},
            ]

        return {"status": "success", "radar_data": radar_data, "bar_data": bar_data, "weakest_subject": weakest_subject,
                "total_students": len(df)}
    except Exception as e:
        db.rollback()
        return JSONResponse(status_code=500, content={"error": str(e)})


@app.get("/api/my_score")
async def get_my_score(account_id: str, db: Session = Depends(get_db)):
    student_record = db.query(StudentScore).filter(StudentScore.user_id == account_id).first()
    if not student_record or not student_record.score_data:
        return {"status": "not_found", "message": "尚未导入成绩"}
    my_radar = [{"subject": k, "A": v, "fullMark": 100} for k, v in student_record.score_data.items()]
    my_weakest = min(my_radar, key=lambda x: x["A"])["subject"] if my_radar else "暂无"
    return {"status": "success", "student_name": student_record.student_name, "radar_data": my_radar,
            "weakest_subject": my_weakest}


@app.post("/api/export_doc")
async def export_doc_endpoint(request: ExportDocRequest):
    try:
        doc = docx.Document()
        doc.add_heading('A25 智能教研产出物', level=0)
        for line in request.content.split('\n'):
            line = line.strip()
            if not line: continue
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            else:
                doc.add_paragraph(line)
        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)
        return StreamingResponse(output_stream,
                                 media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                 headers={'Content-Disposition': 'attachment; filename="A25_Smart_Plan.docx"'})
    except Exception as e:
        raise HTTPException(status_code=500, detail="生成文档失败")


@app.post("/api/upload_knowledge")
async def upload_knowledge_endpoint(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".pdf"): return JSONResponse(status_code=400, content={"error": "仅支持 PDF"})
    try:
        os.makedirs("./docs", exist_ok=True)
        file_path = f"docs/{file.filename}"
        async with aiofiles.open(file_path, 'wb') as out_file:
            await out_file.write(await file.read())
        build_vector_store(file_path)
        return {"status": "success", "message": "知识库入库成功！"}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})


@app.post("/api/generate_practice")
async def generate_practice(request: TargetedPracticeRequest):
    try:
        result_json_str = await agent.generate_targeted_exercises(course_id=request.course_id, error_annotation={
            "explanation": request.error_context, "suggestion": request.suggestion})
        print(f"【DEBUG】大模型原始输出: {result_json_str}")  # 就加这一行！
        return {"status": "success", "data": json.loads(result_json_str)}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.post("/api/enhance_prompt")
async def enhance_prompt(request: PromptEnhanceRequest):
    system_instruction = "你是一位世界顶级的 AI 提示词工程师...直接输出结果。"
    try:
        messages = [SystemMessage(content=system_instruction), HumanMessage(content=f"扩写：{request.short_desc}")]
        response = await agent.llm.ainvoke(messages)
        return {"status": "success", "enhanced_prompt": response.content}
    except Exception as e:
        return {"status": "error", "message": str(e)}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True, reload_excludes=["*.db", "*.db-journal"])

@app.post("/api/register")
async def register(request: RegisterRequest, db: Session = Depends(get_db)):
    # 检查账号是否已存在
    existing = db.query(User).filter(User.account_id == request.account_id).first()
    if existing:
        raise HTTPException(status_code=400, detail="账号已存在")
    
    # 创建新用户
    new_user = User(
        account_id=request.account_id,
        password=request.password,  # 生产环境要哈希存储！
        role=request.role,
        name=request.name
    )
    db.add(new_user)
    db.commit()
    
    return {"status": "success", "message": "注册成功"}